import re, sys, os
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

# dark = band, headings, bullets text accents. light = name/contact on band; panel = light mixed 70% with white.
THEMES = {"navy-ice": dict(label="Navy & Ice", head="Calibri", body="Calibri", dark="14284B", light="CFE0F1")}
def hexrgb(h): return tuple(int(h[i:i+2],16) for i in (0,2,4))
def mix(a,b,t): return ''.join(f'{round(x*(1-t)+y*t):02X}' for x,y in zip(hexrgb(a),hexrgb(b)))
def lum(h):
    def c(v):
        v/=255; return v/12.92 if v<=0.03928 else ((v+0.055)/1.055)**2.4
    r,g,b=[c(x) for x in hexrgb(h)]; return .2126*r+.7152*g+.0722*b
def contrast(a,b):
    la,lb=sorted([lum(a),lum(b)],reverse=True); return (la+.05)/(lb+.05)
for t in THEMES.values(): t['panel']=mix(t['light'],'FFFFFF',0.72)

def rgb(h): return RGBColor.from_string(h)
def font(run, name, size=None, bold=None, color=None, spacing=None):
    run.font.name=name; run._element.rPr.rFonts.set(qn('w:eastAsia'),name)
    if size: run.font.size=Pt(size)
    if bold is not None: run.bold=bold
    if color: run.font.color.rgb=rgb(color)
    if spacing:
        sp=OxmlElement('w:spacing'); sp.set(qn('w:val'),str(spacing)); run._element.rPr.append(sp)
def border(p,color,side='bottom',sz=6,space=2):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr'); e=OxmlElement(f'w:{side}')
    e.set(qn('w:val'),'single'); e.set(qn('w:sz'),str(sz)); e.set(qn('w:space'),str(space)); e.set(qn('w:color'),color); b.append(e); pPr.append(b)
def shade(p,fill,space=8):
    pPr=p._p.get_or_add_pPr(); b=OxmlElement('w:pBdr')
    for side in ('top','left','bottom','right'):
        e=OxmlElement(f'w:{side}'); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'4'); e.set(qn('w:space'),str(space)); e.set(qn('w:color'),fill); b.append(e)
    pPr.append(b)
    sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); pPr.append(sh)
def runs(p,text,t,size,color=None,boldcolor=None):
    for part in re.split(r'(\*\*.+?\*\*)',text):
        if not part: continue
        if part.startswith('**'): font(p.add_run(part[2:-2]),t['body'],size,True,boldcolor or color)
        else: font(p.add_run(part),t['body'],size,False,color)

def base(t,resume):
    d=Document()
    for s in d.sections:
        s.left_margin=s.right_margin=Cm(2.0 if resume else 2.5); s.top_margin=s.bottom_margin=Cm(1.8 if resume else 2.2)
    st=d.styles['Normal']; st.font.name=t['body']; st.font.size=Pt(10.5 if resume else 11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'),t['body'])
    st.paragraph_format.line_spacing=1.15; st.paragraph_format.space_after=Pt(3 if resume else 6)
    return d

def heading(d,t,text):
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(10); p.paragraph_format.space_after=Pt(3); p.paragraph_format.keep_with_next=True
    font(p.add_run(text.upper()),t['head'],10.5,True,t['dark'],spacing=30); border(p,t['light'],'bottom',12,2)

def build_resume(t,src,out):
    d=base(t,True); lines=[l for l in open(src).read().split('\n') if l.strip()]
    while lines and not lines[0].startswith('# '): lines.pop(0)  # drop a sample-resume header block
    name,contact=lines[0][2:],lines[1]
    p=d.add_paragraph(); p.paragraph_format.space_before=Pt(6); p.paragraph_format.space_after=Pt(0)
    font(p.add_run(name),t['head'],24,True,t['light']); shade(p,t['dark'])
    p=d.add_paragraph(); p.paragraph_format.space_after=Pt(12)
    font(p.add_run(contact),t['body'],9.5,False,t['light']); shade(p,t['dark'])
    section=''; titled=False
    for line in lines[2:]:
        if line.startswith('## '):
            h=line[3:]
            if not titled:  # the first ## line is the resume title
                titled=True
                p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0)
                font(p.add_run(h),t['head'],13,True,t['dark'])
            else: section=h; heading(d,t,h)
        elif line.startswith('- '):
            p=d.add_paragraph(style='List Bullet'); runs(p,line[2:],t,10.5); p.paragraph_format.space_after=Pt(1.5)
        elif line.startswith('**') and ' | ' in line and not line.startswith('**Key'):
            p=d.add_paragraph(); p.paragraph_format.space_before=Pt(5); p.paragraph_format.space_after=Pt(1); p.paragraph_format.keep_with_next=True
            bits=line.split(' | '); runs(p,bits[0],t,10.5,t['dark'])
            font(p.add_run('   '+'   ·   '.join(bits[1:])),t['body'],10,False,'5F6670')
        else:
            txt=line.replace(' | ','   ·   ') if section=='Core Capabilities' and not line.startswith('**') else line
            p=d.add_paragraph(); p.paragraph_format.space_after=Pt(0 if section=='Core Capabilities' else 3)
            runs(p,txt,t,10.5,None,t['dark'])
            if section=='Core Capabilities': shade(p,t['panel'],6)
    d.save(out)

def build_letter(t,src,out):
    d=base(t,False)
    for h,fill in ((7,t['dark']),(3,t['light'])):
        p=d.add_paragraph(); pf=p.paragraph_format; pf.space_after=Pt(0); pf.line_spacing_rule=WD_LINE_SPACING.EXACTLY; pf.line_spacing=Pt(h)
        pPr=p._p.get_or_add_pPr(); sh=OxmlElement('w:shd'); sh.set(qn('w:val'),'clear'); sh.set(qn('w:color'),'auto'); sh.set(qn('w:fill'),fill); pPr.append(sh)
    sp=d.add_paragraph(); sp.paragraph_format.space_after=Pt(10)
    paras=open(src).read().strip().split('\n\n')
    for i,para in enumerate(paras):
        p=d.add_paragraph(); p.paragraph_format.space_after=Pt(8)
        for j,ln in enumerate(para.split('\n')):
            if j: p.add_run().add_break()
            last = i==len(paras)-1 and j==len(para.split('\n'))-1
            if last: font(p.add_run(ln),t['head'],12,True,t['dark'])
            else: font(p.add_run(ln),t['body'],11)
    d.save(out)

if __name__=='__main__':
    import argparse
    ap=argparse.ArgumentParser(description="Build Navy & Ice resume and cover letter .docx from markdown")
    ap.add_argument('--resume'); ap.add_argument('--letter')
    ap.add_argument('--out-dir',required=True); ap.add_argument('--prefix',default='Candidate')
    a=ap.parse_args(); os.makedirs(a.out_dir,exist_ok=True); t=THEMES['navy-ice']
    if not (a.resume or a.letter): ap.error('give --resume and/or --letter')
    if a.resume: build_resume(t,a.resume,os.path.join(a.out_dir,f'{a.prefix}-Resume.docx'))
    if a.letter: build_letter(t,a.letter,os.path.join(a.out_dir,f'{a.prefix}-Cover-Letter.docx'))
    print('contrast name-on-band %.1f:1'%contrast(t['light'],t['dark']))
